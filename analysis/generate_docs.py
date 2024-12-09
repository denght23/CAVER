import fitz  # PyMuPDF
from docx import Document
from docx.shared import Inches
import os

def extract_topology_from_pdf(pdf_path):
    """
    从 PDF 文件中提取第一行文本内容作为拓扑设置。
    """
    doc = fitz.open(pdf_path)
    page = doc[0]  # 获取第一页
    text = page.get_text().strip()
    first_line = text.split("\n")[0]  # 提取第一行
    doc.close()
    return first_line

def convert_pdf_to_image(pdf_path, output_folder, zoom_x=2.0, zoom_y=2.0):
    """
    将 PDF 文件转换为高质量图片并保存，返回图片路径。

    :param pdf_path: PDF 文件路径
    :param output_folder: 输出文件夹路径
    :param zoom_x: 水平方向缩放比例 (默认 2.0 为 200%)
    :param zoom_y: 垂直方向缩放比例 (默认 2.0 为 200%)
    :return: 生成图片的路径
    """
    if not os.path.exists(output_folder):
        os.makedirs(output_folder)
    
    doc = fitz.open(pdf_path)
    page = doc[0]  # 仅处理第一页

    # 设置缩放比例以提升分辨率
    mat = fitz.Matrix(zoom_x, zoom_y)
    pix = page.get_pixmap(matrix=mat)  # 应用缩放矩阵
    image_path = os.path.join(output_folder, os.path.basename(pdf_path).replace('.pdf', '.png'))
    pix.save(image_path)

    doc.close()
    return image_path

def generate_word_with_images(pdf_paths, output_docx, output_image_folder):
    """
    将 PDF 文件转换为图片，并根据命名规则生成 Word 文档。
    """
    doc = Document()
    grouped_files = {}
    
    # 分组文件，根据文件名中的 "AVG" 和 "P99"
    for pdf_path in pdf_paths:
        base_name = os.path.basename(pdf_path)
        key = base_name.replace("AVG_", "").replace("P99_", "").replace(".pdf", "")
        if key not in grouped_files:
            grouped_files[key] = {}
        if "AVG" in base_name:
            grouped_files[key]['AVG'] = pdf_path
        elif "P99" in base_name:
            grouped_files[key]['P99'] = pdf_path
    
    # 处理分组后的文件
    for key, files in grouped_files.items():
        # 添加拓扑文字
        if key:
            doc.add_paragraph(key)
        
        # 转换图片并插入到 Word
        table = doc.add_table(rows=1, cols=2)
        row = table.rows[0]
        
        if 'AVG' in files:
            avg_image_path = convert_pdf_to_image(files['AVG'], output_image_folder)
            cell = row.cells[0]
            run = cell.paragraphs[0].add_run()
            run.add_picture(avg_image_path, width=Inches(2.5))
        
        if 'P99' in files:
            p99_image_path = convert_pdf_to_image(files['P99'], output_image_folder)
            cell = row.cells[1]
            run = cell.paragraphs[0].add_run()
            run.add_picture(p99_image_path, width=Inches(2.5))
        
        doc.add_paragraph("")  # 添加空行用于分隔
    
    # 保存 Word 文档
    doc.save(output_docx)

# PDF 文件路径列表
pdf_files = [_.strip() for _ in '''
/home/zj/ns-allinone-3.19/ns-3.19/analysis/figures/AVG_TOPO_fat_k8_100G_OS2_LOAD_40_FC_Lossless.pdf
/home/zj/ns-allinone-3.19/ns-3.19/analysis/figures/P99_TOPO_fat_k8_100G_OS2_LOAD_40_FC_Lossless.pdf
/home/zj/ns-allinone-3.19/ns-3.19/analysis/figures/AVG_TOPO_fat_k8_100G_OS2_LOAD_60_FC_Lossless.pdf
/home/zj/ns-allinone-3.19/ns-3.19/analysis/figures/P99_TOPO_fat_k8_100G_OS2_LOAD_60_FC_Lossless.pdf
/home/zj/ns-allinone-3.19/ns-3.19/analysis/figures/AVG_TOPO_fat_k8_100G_OS2_LOAD_80_FC_Lossless.pdf
/home/zj/ns-allinone-3.19/ns-3.19/analysis/figures/P99_TOPO_fat_k8_100G_OS2_LOAD_80_FC_Lossless.pdf
/home/zj/ns-allinone-3.19/ns-3.19/analysis/figures/AVG_TOPO_fat_k8_100G_bond_OS2_LOAD_40_FC_Lossless.pdf
/home/zj/ns-allinone-3.19/ns-3.19/analysis/figures/P99_TOPO_fat_k8_100G_bond_OS2_LOAD_40_FC_Lossless.pdf
/home/zj/ns-allinone-3.19/ns-3.19/analysis/figures/AVG_TOPO_fat_k8_100G_bond_OS2_LOAD_60_FC_Lossless.pdf
/home/zj/ns-allinone-3.19/ns-3.19/analysis/figures/P99_TOPO_fat_k8_100G_bond_OS2_LOAD_60_FC_Lossless.pdf
/home/zj/ns-allinone-3.19/ns-3.19/analysis/figures/AVG_TOPO_fat_k8_100G_bond_OS2_LOAD_80_FC_Lossless.pdf
/home/zj/ns-allinone-3.19/ns-3.19/analysis/figures/P99_TOPO_fat_k8_100G_bond_OS2_LOAD_80_FC_Lossless.pdf
/home/zj/ns-allinone-3.19/ns-3.19/analysis/figures/AVG_TOPO_fat_k8_100G_OS1_LOAD_40_FC_Lossless.pdf
/home/zj/ns-allinone-3.19/ns-3.19/analysis/figures/P99_TOPO_fat_k8_100G_OS1_LOAD_40_FC_Lossless.pdf
/home/zj/ns-allinone-3.19/ns-3.19/analysis/figures/AVG_TOPO_fat_k8_100G_OS1_LOAD_60_FC_Lossless.pdf
/home/zj/ns-allinone-3.19/ns-3.19/analysis/figures/P99_TOPO_fat_k8_100G_OS1_LOAD_60_FC_Lossless.pdf
/home/zj/ns-allinone-3.19/ns-3.19/analysis/figures/AVG_TOPO_fat_k8_100G_OS1_LOAD_80_FC_Lossless.pdf
/home/zj/ns-allinone-3.19/ns-3.19/analysis/figures/P99_TOPO_fat_k8_100G_OS1_LOAD_80_FC_Lossless.pdf
'''.split('\n') if len(_) != 0]
print(pdf_files)

# 输出路径
output_image_folder = "output_images"  # 存放转换后的图片
output_docx = "data.docx"  # 输出 Word 文件名

# 生成 Word 文档
generate_word_with_images(pdf_files, output_docx, output_image_folder)

print(f"Word 文档已生成：{output_docx}")
